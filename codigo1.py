import os
import socket
import threading
import webbrowser
from datetime import datetime
from flask import Flask, render_template, request, redirect, url_for, flash
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

app = Flask(__name__)
app.secret_key = "encuesta-musica"

ruta_excel = os.path.join(os.path.dirname(os.path.abspath(__file__)), "encuesta.xlsx")

encabezados = ["Fecha", "Nombre", "Edad", "Género favorito", "Artista favorito", "Por qué le gusta"]

generos = ["Rock", "Pop", "Reguetón", "Salsa", "Vallenato", "Electrónica","Rap", "Música popular", "Clásica", "Otro"]
def preparar_excel():
    if os.path.exists(ruta_excel):
        return
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Respuestas"
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="3B2466")
    for letra, ancho in zip("ABCDEF", [18, 28, 8, 20, 28, 45]):
        hoja.column_dimensions[letra].width = ancho
    libro.save(ruta_excel)


def guardar_en_excel(fila):
    preparar_excel()
    libro = load_workbook(ruta_excel)
    libro["Respuestas"].append(fila)
    libro.save(ruta_excel)


def validar(datos):
    errores = {}
    if len(datos["nombre"]) < 3:
        errores["nombre"] = "¿Cómo te llamas?"
    if not datos["edad"].isdigit() or not 5 <= int(datos["edad"]) <= 99:
        errores["edad"] = "Pon tu edad en números."
    if datos["genero"] not in generos:
        errores["genero"] = "Escoge tu género favorito."
    elif datos["genero"] == "Otro" and not datos["otro"]:
        errores["otro"] = "Cuéntanos cuál es."
    return errores


@app.route("/", methods=["GET", "POST"])
def inicio():
    datos = {}
    errores = {}

    if request.method == "POST":
        datos = {
            "nombre": request.form.get("nombre", "").strip(),
            "edad": request.form.get("edad", "").strip(),
            "genero": request.form.get("genero", ""),
            "otro": request.form.get("otro", "").strip(),
            "artista": request.form.get("artista", "").strip(),
            "razon": request.form.get("razon", "").strip(),
        }
        errores = validar(datos)

        if not errores:
            genero = datos["otro"] if datos["genero"] == "Otro" else datos["genero"]
            fila = [
                datetime.now().strftime("%d/%m/%Y %H:%M"),
                datos["nombre"],
                int(datos["edad"]),
                genero,
                datos["artista"],
                datos["razon"],
            ]
            try:
                guardar_en_excel(fila)
                flash(f"¡Gracias, {datos['nombre'].split()[0]}! Tu respuesta quedó guardada.")
                return redirect(url_for("inicio"))
            except PermissionError:
                errores["general"] = "No se pudo guardar porque encuesta.xlsx está abierto. Ciérralo e intenta otra vez."

    return render_template("formulario.html", datos=datos, errores=errores, generos=generos)


def ip_del_computador():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()


if __name__ == "__main__":
    preparar_excel()
    print("Respuestas guardadas en:", ruta_excel)
    print("En este computador: http://127.0.0.1:5000")
    print(f"Desde otro equipo en el mismo wifi: http://{ip_del_computador()}:5000")
    threading.Timer(1, lambda: webbrowser.open("http://127.0.0.1:5000")).start()
    app.run(host="0.0.0.0", port=5000, debug=False)